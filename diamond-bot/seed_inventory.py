"""
Seed the bot with a large diamond inventory from a CSV/XLSX stock file.

These are un-owned "market" listings (no tg_id) tagged with a source group, so
Search Diamond returns real matches out of the box. Natural vs lab-grown is
detected per row and kept in separate matching pools.

    python seed_inventory.py "<path to stock file>" [--source market] [--replace]

--replace first deletes existing seed listings from that source.
"""
from __future__ import annotations

import sys
from pathlib import Path

import db
import ingest


def _rows_from(path: Path):
    data = path.read_bytes()
    name = path.name.lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        import openpyxl, io
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        # tolerate a blank/title first row: pick the first row that looks like headers
        header = None
        for raw in it:
            cells = [str(c).strip() if c is not None else "" for c in raw]
            if sum(bool(c) for c in cells) >= 4:
                header = cells
                break
        if not header:
            return
        for r in it:
            yield dict(zip(header, r))
    else:
        import csv, io
        text = data.decode("utf-8-sig", errors="replace")
        yield from csv.DictReader(io.StringIO(text))


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        raise SystemExit(1)
    path = Path(sys.argv[1])
    source = "market"
    if "--source" in sys.argv:
        source = sys.argv[sys.argv.index("--source") + 1]
    replace = "--replace" in sys.argv

    db.init_db()
    if replace:
        with db._conn() as con:  # noqa: SLF001
            n = con.execute("DELETE FROM listings WHERE source=? AND tg_id IS NULL",
                            (source,)).rowcount
        print(f"removed {n} existing '{source}' seed listings")

    stones, skipped = [], 0
    for row in _rows_from(path):
        norm = ingest._map_row(row)  # noqa: SLF001
        if norm:
            stones.append(norm)
        else:
            skipped += 1
    n = db.add_listings_bulk(stones, tg_id=None, source=source, source_group=source) if stones else 0
    nat = sum(1 for s in stones if "lab_grown" not in (s.get("flags") or []))
    print(f"seeded {n} stones from {path.name} (source={source}); "
          f"{nat} natural / {n - nat} lab-grown; {skipped} skipped")


if __name__ == "__main__":
    main()
